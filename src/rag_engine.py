import os
import time
import logging
import subprocess

# --- CẤU HÌNH GOOGLE CHAT (REST) ---
os.environ["GRPC_VERBOSITY"] = "ERROR"
os.environ["GLOG_minloglevel"] = "2"
import google.generativeai as genai

if os.getenv("GOOGLE_API_KEY"):
    genai.configure(api_key=os.getenv("GOOGLE_API_KEY"), transport="rest")

from langchain_community.document_loaders import TextLoader, PyPDFLoader, Docx2txtLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_cohere import CohereEmbeddings, CohereRerank
from langchain.retrievers import ContextualCompressionRetriever
# --- THƯ VIỆN PINECONE ---
from langchain_pinecone import PineconeVectorStore
from pinecone import Pinecone

logger = logging.getLogger(__name__)

DATA_DIR = "data"


class EnterpriseRAG:
    def __init__(self):
        self.api_key = os.getenv("GOOGLE_API_KEY")
        self.cohere_key = os.getenv("COHERE_API_KEY")
        self.pinecone_api_key = os.getenv("PINECONE_API_KEY")
        self.index_name = os.getenv("PINECONE_INDEX_NAME")

        # Cấu hình Cohere Embeddings (1024 dimensions)
        if self.cohere_key:
            self.embedding_model = CohereEmbeddings(
                cohere_api_key=self.cohere_key,
                model="embed-multilingual-v3.0",
            )
        else:
            self.embedding_model = None

        # Các đối tượng nặng được khởi tạo một lần rồi tái sử dụng
        self._vector_store = None
        self._reranker = None
        self._llm = None
        self._llm_key = None

    # --- Tài nguyên dùng lại (cache) ---
    def _get_vector_store(self):
        if self._vector_store is None:
            self._vector_store = PineconeVectorStore(
                index_name=self.index_name,
                embedding=self.embedding_model,
                pinecone_api_key=self.pinecone_api_key,
            )
        return self._vector_store

    def _get_reranker(self):
        if self._reranker is None and self.cohere_key:
            self._reranker = CohereRerank(
                cohere_api_key=self.cohere_key,
                model="rerank-multilingual-v3.0",
                top_n=5,
            )
        return self._reranker

    def _get_llm(self):
        # Tạo lại nếu API key đổi (cho phép truyền key động qua API/UI)
        if self._llm is None or self._llm_key != self.api_key:
            self._llm = ChatGoogleGenerativeAI(
                model="gemini-2.5-flash",
                google_api_key=self.api_key,
                temperature=0.3,
                transport="rest",
                max_output_tokens=8192,
            )
            self._llm_key = self.api_key
        return self._llm

    # --- Đọc các định dạng đặc biệt ---
    @staticmethod
    def _read_doc(path):
        """Đọc file Word .doc cũ (OLE2) bằng antiword với bảng mã UTF-8."""
        try:
            result = subprocess.run(
                ["antiword", "-m", "UTF-8.txt", path],
                capture_output=True,
                timeout=120,
            )
            return result.stdout.decode("utf-8", errors="ignore")
        except FileNotFoundError:
            logger.warning("Chưa cài 'antiword' — bỏ qua file .doc: %s", path)
        except Exception as e:
            logger.warning("Lỗi đọc .doc %s: %s", path, e)
        return ""

    @staticmethod
    def _read_xlsx(path):
        """Trích xuất văn bản từ tất cả các sheet của file Excel .xlsx."""
        import openpyxl

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        parts.append("\t".join(cells))
            wb.close()
            return "\n".join(parts)
        except Exception as e:
            logger.warning("Lỗi đọc .xlsx %s: %s", path, e)
            return ""

    def _load_documents(self):
        """Quét toàn bộ thư mục data và nạp mọi định dạng được hỗ trợ."""
        documents = []
        for root, _dirs, files in os.walk(DATA_DIR):
            category = os.path.basename(root) if root != DATA_DIR else "General"
            for name in files:
                path = os.path.join(root, name)
                ext = os.path.splitext(name)[1].lower()
                docs = []
                try:
                    if ext == ".txt":
                        docs = TextLoader(path, encoding="utf-8").load()
                    elif ext == ".pdf":
                        docs = PyPDFLoader(path).load()
                    elif ext == ".docx":
                        docs = Docx2txtLoader(path).load()
                    elif ext == ".doc":
                        text = self._read_doc(path)
                        docs = [Document(page_content=text, metadata={"source": path})] if text.strip() else []
                    elif ext in (".xlsx", ".xlsm"):
                        text = self._read_xlsx(path)
                        docs = [Document(page_content=text, metadata={"source": path})] if text.strip() else []
                    else:
                        logger.info("Bỏ qua định dạng chưa hỗ trợ: %s", path)
                        continue
                except Exception as e:
                    logger.warning("Lỗi load %s: %s", path, e)
                    continue

                for doc in docs:
                    doc.metadata["category"] = category
                    doc.metadata["source_name"] = name
                documents.extend(docs)
        return documents

    def index_knowledge_base(self):
        if not self.cohere_key: return "❌ Lỗi: Thiếu COHERE_API_KEY."
        if not self.pinecone_api_key: return "❌ Lỗi: Thiếu PINECONE_API_KEY."
        if not self.index_name: return "❌ Lỗi: Thiếu PINECONE_INDEX_NAME."

        if not os.path.exists(DATA_DIR):
            os.makedirs(DATA_DIR)
            return "Folder data created."

        print("--- 🚀 START INDEXING TO PINECONE CLOUD ---")

        # 1. Quét tài liệu (mọi định dạng)
        all_documents = self._load_documents()
        if not all_documents:
            return "No documents found."

        # 2. Cắt nhỏ văn bản
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
        texts = text_splitter.split_documents(all_documents)
        print(f"Tổng: {len(texts)} đoạn văn.")

        try:
            # 3. Kết nối Pinecone và xóa dữ liệu cũ (làm sạch Index)
            pc = Pinecone(api_key=self.pinecone_api_key)
            index = pc.Index(self.index_name)

            try:
                index.delete(delete_all=True)
                print("🗑️ Đã xóa dữ liệu cũ trên Cloud.")
                time.sleep(2)  # Đợi Pinecone xử lý xóa
            except Exception as e:
                print(f"⚠️ Không thể xóa (có thể Index trống): {e}")

            # 4. Nạp dữ liệu mới (Batching)
            vector_store = self._get_vector_store()

            batch_size = 20
            total_batches = (len(texts) + batch_size - 1) // batch_size

            for i in range(0, len(texts), batch_size):
                batch = texts[i : i + batch_size]
                vector_store.add_documents(batch)
                print(f"☁️ Pinecone Upload: Xong lô {i // batch_size + 1}/{total_batches}")
                time.sleep(0.5)

            return f"✅ Thành công! Đã đẩy {len(texts)} đoạn văn lên Mây (Pinecone)."

        except Exception as e:
            return f"❌ Lỗi Indexing Pinecone: {str(e)}"

    def retrieve_answer(self, query, chat_history="", category=None):
        if not self.api_key: return "Lỗi: Chưa cấu hình Google API Key."
        if not self.index_name: return "Lỗi: Chưa cấu hình Pinecone Index."

        vector_store = self._get_vector_store()
        reranker = self._get_reranker()
        # Lấy nhiều ứng viên hơn khi có rerank để Cohere chọn lọc lại
        fetch_k = 20 if reranker else 5

        def build_retriever(filter_category=None):
            search_kwargs = {"k": fetch_k}
            if filter_category and filter_category != "General":
                search_kwargs["filter"] = {"category": filter_category}
            base = vector_store.as_retriever(search_kwargs=search_kwargs)
            if reranker:
                return ContextualCompressionRetriever(base_compressor=reranker, base_retriever=base)
            return base

        relevant_docs = []
        try:
            if category and category != "General":
                relevant_docs = build_retriever(category).invoke(query)

            # Không lọc được theo category thì tìm trên toàn bộ kho
            if not relevant_docs:
                relevant_docs = build_retriever().invoke(query)

            if not relevant_docs:
                return "Dạ, hiện tại em chưa tìm thấy thông tin này trong hệ thống dữ liệu."

        except Exception as e:
            return f"Lỗi truy vấn Pinecone: {str(e)}"

        formatted_context = ""
        for doc in relevant_docs:
            source = doc.metadata.get("source_name", "Tài liệu nội bộ")
            content = doc.page_content.replace("\n", " ")
            formatted_context += f"--- TÀI LIỆU: {source} ---\nNội dung: {content}\n\n"

        safe_history = chat_history.replace("{", "(").replace("}", ")")

        # --- PROMPT: Yêu cầu trả lời gọn và đầy đủ ---
        prompt = f"""
        VAI TRÒ:
        Bạn là Trợ lý HR ảo của công ty Takagi Việt Nam. Tên bạn là "Trợ lý HR".
        Bạn xưng hô là "em" và gọi người dùng là "anh/chị".
        Tính cách: Tận tâm, nhẹ nhàng, chuyên nghiệp nhưng gần gũi.

        DỮ LIỆU TRA CỨU:
        {formatted_context}

        LỊCH SỬ TRÒ CHUYỆN:
        {safe_history}

        CÂU HỎI MỚI: "{query}"

        YÊU CẦU TRẢ LỜI:
        1. **chính xác và dễ hiểu:** Hãy trả lời chính xác và dễ hiểu, trích dẫn đầy đủ các điều khoản liên quan để nhân viên hiểu rõ.
        2. **Trình bày rõ ràng:** Sử dụng gạch đầu dòng (-) cho các ý chính. Ngôn ngữ mạch lạc rõ ràng.
        3. **Nguồn tài liệu:** Ghi rõ tên văn bản tham khảo ở cuối câu trả lời (Ví dụ: Theo Noi_quy_lao_dong.pdf).
        4. **Thân thiện:** Giữ giọng văn nhẹ nhàng, tận tâm của HR. Cuối cùng, đề xuất thêm gợi ý hoặc hỏi người dùng có cần thêm hỗ trợ nào khác không, "em" sẵn sàng hỗ trợ bất cứ lúc nào.

        BẮT ĐẦU TRẢ LỜI:
        """

        try:
            response = self._get_llm().invoke(prompt)
            return response.content
        except Exception as e:
            return f"Lỗi Gemini: {str(e)}"
